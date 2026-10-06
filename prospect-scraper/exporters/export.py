"""Export dei prospect in CSV (UTF-8) e XLSX (intestazioni leggibili, filtri,
prima riga bloccata, larghezze colonne ragionevoli)."""

from __future__ import annotations

import io
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models.prospect import Prospect
from qualify.models import Qualification

# (campo, intestazione leggibile) nell'ordine di esportazione
COLUMNS: list[tuple[str, str]] = [
    ("company_name", "Company"), ("website", "Website"), ("domain", "Domain"),
    ("country", "Country"), ("region", "Region"), ("city", "City"), ("address", "Address"),
    ("postal_code", "Postal code"), ("category", "Category"), ("phone", "Phone"),
    ("phones", "All phones"), ("email", "Email"), ("emails", "All emails"),
    ("linkedin", "LinkedIn"), ("instagram", "Instagram"), ("facebook", "Facebook"),
    ("youtube", "YouTube"), ("twitter", "X / Twitter"), ("vat_id", "VAT ID"),
    ("description", "Description"), ("source", "Source"), ("source_url", "Source URL"),
    ("search_query", "Search query"), ("status", "Status"), ("error_message", "Error"),
    ("first_seen", "First seen"), ("last_seen", "Last seen"), ("id", "ID"),
]
LABELS = dict(COLUMNS)
URL_LABELS = {"Website", "LinkedIn", "Instagram", "Facebook", "YouTube", "X / Twitter", "Source URL"}

_PHONE_LIKE = re.compile(r"^\+[\d\s().;/-]+$")


# Colonne della qualifica agenzie: (intestazione, funzione). Le prime vanno dopo "Company",
# le altre in fondo (solo se ``qualifications`` è passato).
def _join(values) -> str:
    return ", ".join(str(v) for v in values or [])


QUAL_FRONT: list[tuple[str, object]] = [
    ("Score", lambda q: q.score), ("Is agency", lambda q: q.is_agency),
    ("SEO level", lambda q: q.seo_level), ("Servizi ricorrenti", lambda q: q.servizi_ricorrenti),
    ("Size", lambda q: q.size_signal), ("Blog", lambda q: q.blog_status),
    ("Ultimo post blog", lambda q: q.blog_last_post), ("Servizi", lambda q: _join(q.servizi)),
    ("Servizi (altro)", lambda q: q.servizi_altro), ("Verticali", lambda q: _join(q.verticali)),
    ("Note", lambda q: q.note),
]
QUAL_BACK: list[tuple[str, object]] = [
    ("Qualifica status", lambda q: q.status), ("Qualifica errore", lambda q: q.error),
    ("Prova agenzia", lambda q: q.evidence.get("agenzia", "")),
    ("Prova SEO", lambda q: q.evidence.get("seo", "")),
    ("Prova ricorrenti", lambda q: q.evidence.get("ricorrenti", "")),
    ("Prova team", lambda q: q.evidence.get("team", "")),
    ("Pagine analizzate", lambda q: _join(q.pages_used)),
    ("Modello AI", lambda q: q.llm_model),
    ("Costo AI ($)", lambda q: round(q.cost_usd, 5) if q.llm_model else None),
]


def prospects_to_dataframe(prospects: list[Prospect],
                           qualifications: dict[int, Qualification] | None = None) -> pd.DataFrame:
    rows = [{label: getattr(p, field) or "" for field, label in COLUMNS} for p in prospects]
    headers = [label for _, label in COLUMNS]
    if qualifications is not None:
        extra = QUAL_FRONT + QUAL_BACK
        for row, p in zip(rows, prospects):
            q = qualifications.get(p.id)
            for label, getter in extra:
                row[label] = getter(q) if q is not None else None
        headers = (headers[:1] + [h for h, _ in QUAL_FRONT] + headers[1:] + [h for h, _ in QUAL_BACK])
    df = pd.DataFrame(rows, columns=headers)
    df["ID"] = pd.to_numeric(df["ID"], errors="coerce").astype("Int64")
    if qualifications is not None:
        df["Score"] = pd.to_numeric(df["Score"], errors="coerce").astype("Int64")
        df["Costo AI ($)"] = pd.to_numeric(df["Costo AI ($)"], errors="coerce")
        for label, _ in QUAL_FRONT + QUAL_BACK:       # niente None/NaN nei testi: celle vuote
            if label not in ("Score", "Costo AI ($)"):
                df[label] = df[label].fillna("")
    return df


def _safe_cell(value):
    """Evita che un testo raccolto dal web venga interpretato come formula (CSV injection)."""
    if not isinstance(value, str) or not value:
        return value
    if value[0] in ("=", "@", "\t", "\r") or (value[0] in "+-" and not _PHONE_LIKE.match(value)):
        return "'" + value
    return value


def _sanitize(df: pd.DataFrame) -> pd.DataFrame:
    # vale sia per dtype object (pandas 2) sia per dtype str (pandas 3)
    return df.apply(lambda col: col if pd.api.types.is_numeric_dtype(col) else col.map(_safe_cell))


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    # UTF-8 con BOM: Excel riconosce correttamente accenti e caratteri speciali.
    return _sanitize(df).to_csv(index=False).encode("utf-8-sig")


def to_xlsx_bytes(df: pd.DataFrame, sheet_name: str = "Prospects") -> bytes:
    return to_xlsx_multi_bytes({sheet_name: df})


def to_xlsx_multi_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Un file XLSX con un foglio per ogni DataFrame (es. un foglio per ricerca)."""
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for name, df in sheets.items():
        _write_sheet(wb.create_sheet(_sheet_title(name, used)), _sanitize(df))
    if not wb.worksheets:
        _write_sheet(wb.create_sheet("Prospects"), _sanitize(pd.DataFrame(columns=list(LABELS.values()))))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_title(name: str, used: set[str]) -> str:
    """Nome foglio valido per Excel: max 31 caratteri, senza []:*?/\\ e unico."""
    base = re.sub(r"[\[\]:*?/\\]", " ", name or "Prospects").strip()[:31] or "Prospects"
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(title.lower())
    return title


def _write_sheet(ws, df: pd.DataFrame) -> None:
    headers = list(df.columns)
    ws.append(headers)
    for row in df.itertuples(index=False):
        ws.append([None if (v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NA) else v
                   for v in row])

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    link_font = Font(color="0563C1", underline="single")
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        values = [str(v) for v in df[header].tolist() if v is not None and not pd.isna(v) and v != ""]
        lengths = sorted(len(v) for v in values)
        typical = lengths[int(len(lengths) * 0.9)] if lengths else 0
        ws.column_dimensions[letter].width = max(10, min(50, max(len(header), typical) + 2))
        if header in URL_LABELS:
            for cell in ws[letter][1:]:
                if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                    cell.hyperlink = cell.value
                    cell.font = link_font

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def slugify(text: str, max_len: int = 50) -> str:
    """Testo -> parte di nome file ("hotel 4 stelle | Palermo" -> "hotel-4-stelle-palermo")."""
    import unicodedata

    text = unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text[:max_len].rstrip("-") or "prospects"
