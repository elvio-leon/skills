"""Ricerche salvate, export multi-foglio e launcher."""

import io
import json

from openpyxl import load_workbook

import launcher
from database.db import Database
from exporters.export import prospects_to_dataframe, slugify, to_xlsx_multi_bytes
from models.prospect import Prospect


def test_list_runs_with_counts(tmp_path):
    db = Database(tmp_path / "r.db")
    r1 = db.start_run(mode="Search", keyword="hotel", location="Palermo")
    r2 = db.start_run(mode="Search", keyword="saas", location="Milano")
    db.start_run(mode="Search", keyword="vuota")                   # nessun prospect: non elencata
    for i in range(3):
        db.link_run(r1, db.save(Prospect(company_name=f"H{i}", website=f"https://h{i}.it")))
    db.link_run(r2, db.save(Prospect(company_name="S", website="https://s.it")))
    runs = db.list_runs()
    assert [(r["id"], r["n_prospects"]) for r in runs] == [(r2, 1), (r1, 3)]  # più recente prima


def test_multi_sheet_export():
    df = prospects_to_dataframe([Prospect(company_name="A", website="https://a.it")])
    data = to_xlsx_multi_bytes({"hotel Palermo": df, "hotel Palermo ": df, "x/y:z" * 10: df})
    wb = load_workbook(io.BytesIO(data))
    assert [ws.title for ws in wb][:2] == ["hotel Palermo", "hotel Palermo (2)"]
    assert all(len(ws.title) <= 31 and "/" not in ws.title for ws in wb)
    assert all(ws.freeze_panes == "A2" and ws.auto_filter.ref for ws in wb)
    assert load_workbook(io.BytesIO(to_xlsx_multi_bytes({}))).active.title == "Prospects"


def test_slugify():
    assert slugify("hotel 4 stelle Palermo") == "hotel-4-stelle-palermo"
    assert slugify("Città d'arte / Più") == "citta-d-arte-piu"
    assert slugify("") == "prospects"


def test_launcher_port_and_state(tmp_path, monkeypatch):
    port = launcher.free_port(0)
    assert isinstance(port, int) and port > 0
    assert not launcher.is_running(port)            # nessun server su quella porta
    monkeypatch.setattr(launcher, "STATE_FILE", tmp_path / "state.json")
    assert launcher.running_port() is None           # file assente
    (tmp_path / "state.json").write_text(json.dumps({"port": port, "pid": 1}))
    assert launcher.running_port() is None           # file presente ma server spento


def test_launcher_skips_install_when_up_to_date(tmp_path, monkeypatch):
    import hashlib

    req = tmp_path / "requirements.txt"
    req.write_text("streamlit\n")
    venv = tmp_path / ".venv"
    py = venv / ("Scripts/python.exe" if launcher.IS_WINDOWS else "bin/python")
    py.parent.mkdir(parents=True)
    py.write_text("")
    (venv / ".requirements.sha256").write_text(hashlib.sha256(req.read_bytes()).hexdigest())
    monkeypatch.setattr(launcher, "REQUIREMENTS", req)
    monkeypatch.setattr(launcher, "VENV", venv)
    calls = []
    monkeypatch.setattr(launcher.subprocess, "check_call", lambda *a, **k: calls.append(a))
    launcher.ensure_environment()
    assert calls == []                                # già installato: nessun pip
    req.write_text("streamlit\npandas\n")             # dipendenze cambiate: reinstalla
    launcher.ensure_environment()
    assert any("pip" in c[0] for c in calls)
