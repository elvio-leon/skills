"""Test end-to-end: ricerca (SearXNG finto) -> dedup -> DB -> enrichment -> export."""

import csv
import io

import pytest
from openpyxl import load_workbook

from config import settings
from core.pipeline import MODE_BOTH, MODE_ENRICH, MODE_SEARCH, Pipeline, RunParams, result_to_prospect
from database.db import Database
from exporters.export import prospects_to_dataframe, to_csv_bytes, to_xlsx_bytes
from scrapers.http import HttpClient
from scrapers.search.base import SearchResult
from scrapers.website import WebsiteCrawler
from tests.site_server import SiteServer


@pytest.fixture(scope="module")
def srv():
    with SiteServer() as s:
        yield s


@pytest.fixture()
def env(srv, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "SEARXNG_URL", f"http://127.0.0.1:{srv.port}")
    monkeypatch.setattr(settings, "USE_PLAYWRIGHT", False)
    client = HttpClient(delay=0, timeout=1, retries=0)
    db = Database(tmp_path / "t.db")
    messages, progress = [], []
    pipe = Pipeline(db, client, WebsiteCrawler(client, time_budget=20),
                    on_message=lambda t, lvl="info": messages.append((lvl, t)),
                    on_progress=lambda d, n, label="": progress.append((d, n)))
    return pipe, db, messages, progress


def test_search_plus_enrichment(env):
    pipe, db, messages, progress = env
    params = RunParams(category="Custom", keyword="hotel", location="Palermo", max_results=20,
                       mode=MODE_BOTH, providers=["searxng"])
    res = pipe.run(params)
    assert not res.fatal_error, res.log_text
    texts = [t for _, t in messages]
    assert texts[0] == "Searching..." and "Deduplicating..." in texts and "Enriching websites..." in texts
    assert res.n_found == 7        # booking.com e il 2° risultato di 127.0.0.2 filtrati dal provider
    assert res.n_unique == 7
    assert any("7 → 7" in t for t in texts)
    assert progress[-1] == (7, 7)
    assert res.n_enriched == 3 and res.n_failed == 4

    rows = {p.domain: p for p in db.list_prospects(ids=res.prospect_ids)}
    alfa = rows["127.0.0.2"]
    assert alfa.status == "enriched" and alfa.email == "info@hotelalfa.it"
    assert alfa.phone == "+39 091 123 4567" and alfa.city == "Palermo"
    assert alfa.linkedin == "https://www.linkedin.com/company/hotel-alfa"
    assert alfa.raw_data["enrichment"]["pages_visited"]
    failed = {d: p.error_message for d, p in rows.items() if p.status == "failed"}
    assert "robots" in failed["127.0.0.3"] and "CAPTCHA" in failed["127.0.0.4"]
    assert "timeout" in failed["127.0.0.5"] and "404" in failed["127.0.0.9"]

    # seconda esecuzione: nessun duplicato, i siti arricchiti vengono dalla cache
    messages.clear()
    res2 = pipe.run(params)
    assert db.count() == 7
    assert res2.n_cached == 3
    assert sorted(res2.prospect_ids) == sorted(res.prospect_ids)

    # export
    df = prospects_to_dataframe(db.list_prospects(ids=res.prospect_ids))
    content = to_csv_bytes(df)
    assert content.startswith(b"\xef\xbb\xbf")
    reader = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    assert len(reader) == 7 and {"Company", "Website", "Email", "Phone", "LinkedIn", "Source",
                                 "Status"} <= set(reader[0])
    wb = load_workbook(io.BytesIO(to_xlsx_bytes(df)))
    ws = wb.active
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
    assert ws["A1"].value == "Company" and ws["A1"].font.bold
    assert all(10 <= ws.column_dimensions[c].width <= 50 for c in ("A", "B", "L", "T"))


def test_search_only_and_enrich_mode(env, srv):
    pipe, db, messages, _ = env
    res = pipe.run(RunParams(keyword="hotel", location="Palermo", max_results=3, mode=MODE_SEARCH,
                             providers=["searxng"]))
    assert res.n_unique == 3 and res.n_enriched == 0
    assert {p.status for p in db.list_prospects(ids=res.prospect_ids)} == {"found"}
    # enrichment dei prospect "found" salvati nel DB
    res2 = pipe.run(RunParams(max_results=10, mode=MODE_ENRICH))
    assert res2.n_unique == 3 and res2.n_enriched + res2.n_failed == 3
    # enrichment di URL inseriti a mano
    res3 = pipe.run(RunParams(max_results=10, mode=MODE_ENRICH, force_refresh=True,
                              urls=[srv.url("127.0.0.7", "/qualsiasi"), "https://www.facebook.com/x"]))
    assert res3.n_unique == 1 and res3.n_enriched == 1


def test_result_mapping_never_invents_website():
    params = RunParams(category="Hospitality", keyword="hotel", location="Palermo")
    social = result_to_prospect(SearchResult("Hotel Social", "https://www.facebook.com/hotelsocial/", "",
                                             "osm", {"company_name": "Hotel Social", "city": "Palermo"}), params)
    assert social.website == "" and social.domain is None
    assert social.facebook == "https://www.facebook.com/hotelsocial" and social.status == "no_website"
    portal = result_to_prospect(SearchResult("X", "https://www.booking.com/hotel/x", "", "searxng"), params)
    assert portal.website == "" and portal.raw_data["non_company_url"]
    ok = result_to_prospect(SearchResult("Home | Hotel Gamma", "https://www.hotelgamma.it/camere?x=1", "",
                                         "searxng", {"phone": "091 111 2222", "instagram": "hotelgamma"}), params)
    assert ok.website == "https://www.hotelgamma.it/" and ok.domain == "hotelgamma.it"
    assert ok.company_name == "Hotel Gamma" and ok.phone == "+39 091 111 2222"
    assert ok.instagram == "https://www.instagram.com/hotelgamma"


def test_name_only_prospect_is_promoted_when_domain_found(tmp_path):
    db = Database(tmp_path / "p.db")
    pipe = Pipeline(db)
    from models.prospect import Prospect
    run = db.start_run(mode="Search")
    first = pipe._save_merged(Prospect(company_name="Hotel Delta", city="Palermo", status="no_website"), run)
    second = pipe._save_merged(Prospect(company_name="Hotel Delta", city="Palermo",
                                        website="https://hoteldelta.it", status="found"), run)
    assert first.id == second.id and db.count() == 1
    assert db.get(first.id).domain == "hoteldelta.it" and db.get(first.id).status == "found"


def test_csv_injection_is_neutralized():
    from exporters.export import _safe_cell
    assert _safe_cell("=HYPERLINK(\"x\")") == "'=HYPERLINK(\"x\")"
    assert _safe_cell("+39 091 123 4567") == "+39 091 123 4567"
    assert _safe_cell("-50% sconto") == "'-50% sconto"


def test_csv_and_xlsx_exports_are_sanitized():
    from models.prospect import Prospect
    df = prospects_to_dataframe([Prospect(company_name='=HYPERLINK("http://x")', phone="+39 091 123 4567")])
    row = list(csv.reader(io.StringIO(to_csv_bytes(df).decode("utf-8-sig"))))[1]
    assert row[0].startswith("'=") and "+39 091 123 4567" in row
    ws = load_workbook(io.BytesIO(to_xlsx_bytes(df))).active
    assert ws["A2"].value.startswith("'=") and ws["A2"].data_type == "s"
