"""Strumenti Python di Metatropi, eseguiti dal Python portatile incluso nell'app.

Uso:
  metatropi_tools.py docx   <in.pdf>  <out.docx>   PDF -> Word (pdf2docx)
  metatropi_tools.py tables <in.pdf>  <out.json>   tabelle del PDF in JSON
  metatropi_tools.py info   <in.pdf>               pagine e caratteri di testo (JSON su stdout)
  metatropi_tools.py pptx   <in.pptx> <out_dir>    slide -> Markdown + immagini (stampa il percorso del .md)

L'avanzamento viene scritto su stderr come righe "PROGRESS <fatto> <totale>".
"""

import json
import logging
import os
import posixpath
import re
import sys
import warnings
import zipfile
from xml.etree import ElementTree as ET

warnings.filterwarnings("ignore")


def progress(done, total):
    sys.stderr.write(f"PROGRESS {done} {total}\n")
    sys.stderr.flush()


# ---------------------------------------------------------------- PDF -> Word

class _ProgressHandler(logging.Handler):
    """pdf2docx registra righe come "(3/10) Page 3": le trasformiamo in avanzamento."""

    def __init__(self, phases):
        super().__init__()
        self.phase = 0
        self.phases = phases

    def emit(self, record):
        msg = record.getMessage()
        if msg.startswith("[") and "Parsing pages" in msg:
            self.phase = 0
        elif "Creating pages" in msg:
            self.phase = 1
        m = re.search(r"\((\d+)/(\d+)\)", msg)
        if m:
            done, total = int(m.group(1)), int(m.group(2))
            progress(self.phase * total + done, self.phases * total)


def cmd_docx(src, dst):
    from pdf2docx import Converter

    logging.getLogger().handlers.clear()
    handler = _ProgressHandler(phases=2)
    logging.getLogger().addHandler(handler)
    logging.getLogger().setLevel(logging.INFO)
    cv = Converter(src)
    try:
        cv.convert(dst, multi_processing=False)
    finally:
        cv.close()


# ---------------------------------------------------------------- tabelle

def _clean(cell):
    if cell is None:
        return ""
    if isinstance(cell, str):
        return re.sub(r"\s*\n\s*", " ", cell).strip()
    return cell


def _docx_tables(docx_path):
    import docx

    document = docx.Document(docx_path)
    found = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells, seen = [], set()
            for cell in row.cells:
                # le celle unite compaiono più volte: si tengono una volta sola
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                cells.append(_clean(cell.text))
            if any(cells):
                rows.append(cells)
        if rows:
            found.append({"index": len(found) + 1, "rows": rows})
    return found


def cmd_tables(src, dst):
    """Le tabelle vengono riconosciute da pdf2docx (anche quelle senza bordi);
    se non ne trova, si prova il riconoscimento per linee di PyMuPDF."""
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        docx_path = os.path.join(tmp, "tabelle.docx")
        cmd_docx(src, docx_path)
        found = _docx_tables(docx_path)
    if not found:
        import pymupdf

        doc = pymupdf.open(src)
        for page in doc:
            try:
                tabs = page.find_tables()
            except Exception:
                continue
            for table in tabs.tables:
                rows = [[_clean(c) for c in row] for row in table.extract()]
                rows = [r for r in rows if any(str(c).strip() for c in r)]
                if rows:
                    found.append({"index": len(found) + 1, "rows": rows})
    with open(dst, "w", encoding="utf-8") as fh:
        json.dump(found, fh, ensure_ascii=False)


def cmd_info(src):
    import pymupdf

    doc = pymupdf.open(src)
    chars = sum(len(page.get_text("text").strip()) for page in doc)
    print(json.dumps({"pages": doc.page_count, "chars": chars}))


# ---------------------------------------------------------------- PowerPoint

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
R_ID = "{%s}id" % NS["r"]
R_EMBED = "{%s}embed" % NS["r"]


def _rels(zf, part):
    folder, name = posixpath.split(part)
    rels_path = posixpath.join(folder, "_rels", name + ".rels")
    if rels_path not in zf.namelist():
        return {}
    root = ET.fromstring(zf.read(rels_path))
    out = {}
    for rel in root.findall("rel:Relationship", NS):
        target = rel.get("Target")
        if rel.get("TargetMode") == "External":
            continue
        out[rel.get("Id")] = posixpath.normpath(posixpath.join(folder, target))
    return out


def _md_escape(text):
    return re.sub(r"([\\`*_\[\]#<>|])", r"\\\1", text)


def _paragraph_text(p):
    parts = []
    for node in p.iter():
        tag = node.tag.split("}")[-1]
        if tag == "t" and node.text:
            parts.append(node.text)
        elif tag == "br":
            parts.append(" ")
    return "".join(parts).strip()


def _shape_md(shape, rels, zf, media_dir, media_names, lines, title_holder):
    tag = shape.tag.split("}")[-1]
    if tag == "grpSp":
        for child in shape:
            _shape_md(child, rels, zf, media_dir, media_names, lines, title_holder)
        return
    if tag == "sp":
        ph = shape.find(".//p:nvPr/p:ph", NS)
        is_title = ph is not None and ph.get("type") in ("title", "ctrTitle")
        body = shape.find("p:txBody", NS)
        if body is None:
            return
        paragraphs = []
        for p in body.findall("a:p", NS):
            text = _paragraph_text(p)
            if not text:
                continue
            ppr = p.find("a:pPr", NS)
            level = int(ppr.get("lvl", "0")) if ppr is not None else 0
            bullet = ppr is not None and ppr.find("a:buNone", NS) is None and (ph is None or ph.get("type") in (None, "body", "obj"))
            paragraphs.append((text, level, bullet))
        if not paragraphs:
            return
        if is_title and title_holder[0] is None:
            title_holder[0] = " ".join(t for t, _, _ in paragraphs)
            return
        for text, level, bullet in paragraphs:
            if bullet or ph is not None and ph.get("type") in (None, "body", "obj"):
                lines.append("    " * level + "- " + _md_escape(text))
            else:
                lines.append(_md_escape(text))
                lines.append("")
        lines.append("")
        return
    if tag == "pic":
        blip = shape.find(".//a:blip", NS)
        target = rels.get(blip.get(R_EMBED)) if blip is not None else None
        if target and target in zf.namelist():
            ext = posixpath.splitext(target)[1].lower()
            if ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".svg"):
                name = f"immagine-{len(media_names) + 1}{ext}"
                with open(os.path.join(media_dir, name), "wb") as fh:
                    fh.write(zf.read(target))
                media_names.append(name)
                lines.append(f"![]({name})")
                lines.append("")
        return
    if tag == "graphicFrame":
        tbl = shape.find(".//a:tbl", NS)
        if tbl is None:
            return
        rows = []
        for tr in tbl.findall("a:tr", NS):
            cells = []
            for tc in tr.findall("a:tc", NS):
                text = " ".join(_paragraph_text(p) for p in tc.findall(".//a:p", NS)).strip()
                cells.append(_md_escape(text))
            rows.append(cells)
        if rows:
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("|" + " --- |" * width)
            for r in rows[1:]:
                lines.append("| " + " | ".join(r) + " |")
            lines.append("")


def cmd_pptx(src, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    zf = zipfile.ZipFile(src)
    pres = "ppt/presentation.xml"
    pres_rels = _rels(zf, pres)
    root = ET.fromstring(zf.read(pres))
    slides = [pres_rels.get(s.get(R_ID)) for s in root.findall("p:sldIdLst/p:sldId", NS)]
    slides = [s for s in slides if s and s in zf.namelist()]
    md, media_names = [], []
    for number, slide in enumerate(slides, 1):
        rels = _rels(zf, slide)
        tree = ET.fromstring(zf.read(slide))
        sp_tree = tree.find("p:cSld/p:spTree", NS)
        lines, title = [], [None]
        if sp_tree is not None:
            for shape in sp_tree:
                _shape_md(shape, rels, zf, out_dir, media_names, lines, title)
        md.append(f"# {_md_escape(title[0]) if title[0] else f'Slide {number}'}")
        md.append("")
        md.extend(lines)
        progress(number, len(slides))
    path = os.path.join(out_dir, "presentazione.md")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(md).rstrip() + "\n")
    print(path)


def main(argv):
    if len(argv) < 2:
        sys.stderr.write(__doc__)
        return 2
    cmd, rest = argv[1], argv[2:]
    commands = {"docx": cmd_docx, "tables": cmd_tables, "info": cmd_info, "pptx": cmd_pptx}
    if cmd not in commands:
        sys.stderr.write(f"Comando sconosciuto: {cmd}\n")
        return 2
    commands[cmd](*rest)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
